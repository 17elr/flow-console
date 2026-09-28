from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import PurePosixPath

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .models import AssetVersion, ImageBatch, ImageJob, ProductMaster, SkuAsset, SkuVariant
from .storage import AssetValidationError, configured_storage, validate_image


CATEGORY = "服装、鞋靴和珠宝饰品 > 女士时尚 > 女士饰品 > 女士项链 > 女士时尚吊坠项链"
ALIEXPRESS_CATEGORY = "珠宝饰品及配件 (Jewelry & Accessories)/流行饰品 (Fashion Jewelry)/项链 (Necklace)"
MAIN_ROLES = (
    "SPU_WHITE_MAIN",
    "SPU_DETAIL_1",
    "SPU_DETAIL_2",
    "SPU_SIZE_INFO",
    "SCENE_MODEL_WEAR",
    "SCENE_LIFESTYLE",
)
ROLE_ALIASES = {
    "SPU_WHITE_MAIN": ("01", "1", "主图", "白底主图", "main", "white_main", "white-main"),
    "SPU_DETAIL_1": ("02", "2", "细节1", "细节图1", "detail1", "detail_1"),
    "SPU_DETAIL_2": ("03", "3", "细节2", "细节图2", "detail2", "detail_2"),
    "SPU_SIZE_INFO": ("04", "4", "尺寸", "尺寸图", "size", "size_info"),
    "SCENE_MODEL_WEAR": ("05", "5", "模特", "佩戴", "model", "model_wear"),
    "SCENE_LIFESTYLE": ("06", "6", "场景", "生活场景", "lifestyle", "scene"),
}


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _number(value: object, default: float = 0) -> float:
    try:
        return float(value) if value not in (None, "") else default
    except (TypeError, ValueError):
        return default


def _first_value(row: dict[str, object], *names: str) -> object:
    for name in names:
        value = row.get(name)
        if value not in (None, ""):
            return value
    return None


def _weight_in_grams(row: dict[str, object], kg_names: tuple[str, ...], gram_names: tuple[str, ...]) -> float:
    kilograms = _first_value(row, *(name for kg_name in kg_names for name in (kg_name, kg_name.replace(" KG(批量)", " KG"))))
    if kilograms not in (None, ""):
        return _number(kilograms) * 1000
    return _number(_first_value(row, *gram_names))


def _normalized_product_parameters(row: dict[str, object]) -> dict[str, object]:
    parameters = dict(row)
    normalized_aliases = {"cid", "妙手类目id", "类目id", "temu类目id", "妙手类目编号"}
    for key, raw_value in row.items():
        normalized_key = re.sub(r"[\s_（）()]+", "", _text(key)).casefold()
        if normalized_key in normalized_aliases:
            value = _text(raw_value)
            if value:
                parameters["cid"] = value
                break
    normalized_values = {
        re.sub(r"[\s_（）()]+", "", _text(key)).casefold(): _text(value)
        for key, value in row.items()
    }
    package_type = (
        normalized_values.get("外包装类型")
        or normalized_values.get("outerpackagetype")
        or normalized_values.get("包装类型")
        or ""
    )
    package_type_norm = _norm(package_type)
    if package_type_norm:
        type_map = {
            "硬包装": 0,
            "软包装硬物": 1,
            "软包装软物": 2,
        }
        for label, code in type_map.items():
            if label in package_type_norm:
                parameters["outerPackageType"] = code
                parameters["外包装类型"] = package_type
                break
    package_shape = (
        normalized_values.get("外包装形状")
        or normalized_values.get("outerpackageshape")
        or normalized_values.get("包装形状")
        or ""
    )
    package_shape_norm = _norm(package_shape)
    if package_shape_norm:
        shape_map = {
            "长方体": 1,
            "方体": 2,
        }
        for label, code in shape_map.items():
            if label in package_shape_norm:
                parameters["outerPackageShape"] = code
                parameters["外包装形状"] = package_shape
                break
    individually_packed = (
        normalized_values.get("sku是否独立包装")
        or normalized_values.get("是否独立包装")
        or normalized_values.get("individuallypacked")
        or ""
    )
    individually_packed_norm = _norm(individually_packed)
    if individually_packed_norm:
        packed_map = {
            "是独立包装": 1,
            "不是独立包装": 0,
        }
        for label, code in packed_map.items():
            if individually_packed_norm == _norm(label):
                parameters["individuallyPacked"] = code
                parameters["SKU是否独立包装"] = label
                break
    return parameters


def _rows_dimensionless(sheet) -> list[dict[str, object]]:
    values = list(sheet.iter_rows(values_only=True))
    header_index = next((index for index, row in enumerate(values[:12]) if any(_text(value) == "浜у搧缂栧彿" for value in row)), None)
    if header_index is None:
        header_index = next(
            (index for index, row in enumerate(values[:12])
             if len([value for value in row if _text(value)]) >= 5 and len(_text(row[0])) <= 40),
            None,
        )
    if header_index is None:
        raise ValueError("工作表缺少产品编号表头")
    headers = [_text(value) for value in values[header_index]]
    return [{headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]}
            for row in values[header_index + 1:] if any(_text(value) for value in row) and _text(row[0])]


def _rows(sheet) -> list[dict[str, object]]:
    if sheet.max_row is None:
        return _rows_dimensionless(sheet)
    header_row = next((row for row in range(1, min(sheet.max_row, 12) + 1) if any(_text(cell.value) == "产品编号" for cell in sheet[row])), None)
    if not header_row:
        raise ValueError(f"工作表“{sheet.title}”缺少产品编号表头")
    headers = [_text(cell.value) for cell in sheet[header_row]]
    result = []
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        row = {headers[index]: value for index, value in enumerate(values) if index < len(headers) and headers[index]}
        if _text(row.get("产品编号")):
            result.append(row)
    return result


def parse_finished_workbook(content: bytes) -> tuple[list[dict], list[dict], bool]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    if "商品参数" not in workbook.sheetnames:
        raise ValueError("Excel 必须包含“商品参数”工作表")
    products = _rows(workbook["商品参数"])
    legacy = "SKU规格" in workbook.sheetnames
    skus = _rows(workbook["SKU规格"]) if legacy else []
    if not products:
        raise ValueError("商品参数表没有可导入的产品")
    codes = [_text(row["产品编号"]) for row in products]
    if len(codes) != len(set(code.casefold() for code in codes)):
        raise ValueError("商品参数表存在重复的产品编号")
    return products, skus, legacy


def _folder_for(relative_path: str) -> str | None:
    parts = [part for part in PurePosixPath(relative_path.replace("\\", "/")).parts if part not in {".", ""}]
    if len(parts) >= 3 and parts[-2].casefold() == "sku":
        return parts[-3]
    return parts[-2] if len(parts) >= 2 else None


def _is_sku_path(relative_path: str) -> bool:
    parts = [part for part in PurePosixPath(relative_path.replace("\\", "/")).parts if part not in {".", ""}]
    return len(parts) >= 3 and parts[-2].casefold() == "sku"


def _norm(value: str) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", value.casefold())


def _pick_role(stem: str) -> str | None:
    normalized = _norm(stem)
    for role, aliases in ROLE_ALIASES.items():
        for alias in aliases:
            normalized_alias = _norm(alias)
            if normalized == normalized_alias:
                return role
            if not normalized_alias.isdigit() and normalized_alias in normalized:
                return role
    return None


def _save_finished_asset(db: Session, product: ProductMaster, batch: ImageBatch, role: str, content: bytes, filename: str, sku: SkuVariant | None = None) -> AssetVersion:
    meta = validate_image(content)
    version = (db.scalar(select(func.count(AssetVersion.id)).where(AssetVersion.product_id == product.id, AssetVersion.role == role)) or 0) + 1
    safe_name = re.sub(r"[^0-9A-Za-z._-]+", "-", filename).strip("-") or f"{role}.{meta['extension']}"
    key = f"finished/{product.spu_code}/{batch.id}/{role}-{version}-{safe_name}"
    configured_storage().put(key, content, meta["mime_type"])
    qc = {"status": "MANUAL_FINISHED_UPLOAD", "human_review_required": True, "collage_detected": None}
    asset = AssetVersion(
        product_id=product.id,
        batch_id=batch.id,
        role=role,
        version=version,
        storage_key=key,
        sha256=meta["sha256"],
        mime_type=meta["mime_type"],
        width=meta["width"],
        height=meta["height"],
        byte_size=meta["byte_size"],
        pipeline_version="finished-upload-v1",
        provider="MANUAL_UPLOAD",
        model=None,
        prompt=None,
        qc_json=json.dumps(qc, ensure_ascii=False),
    )
    db.add(asset)
    db.flush()
    job = ImageJob(
        batch_id=batch.id,
        product_id=product.id,
        sku_id=sku.id if sku else None,
        job_type=role,
        stage="COMPLETED",
        status="COMPLETED" if role not in {"SCENE_MODEL_WEAR", "SCENE_LIFESTYLE"} else "READY_FOR_REVIEW",
        retry_count=0,
        max_retries=0,
        result_asset_id=asset.id,
        qc_json=asset.qc_json,
        completed_at=datetime.now(timezone.utc),
    )
    db.add(job)
    if sku:
        db.add(SkuAsset(sku_id=sku.id, batch_id=batch.id, asset_version_id=asset.id, sku_code=sku.sku_code, component_manifest=json.dumps({"source": "finished_upload"})))
    return asset


def import_finished_package(db: Session, workbook_content: bytes, uploads: list[tuple[str, str, bytes]]) -> dict:
    product_rows, sku_rows, legacy = parse_finished_workbook(workbook_content)
    images_by_folder: dict[str, list[tuple[str, bytes, bool]]] = {}
    for filename, relative_path, content in uploads:
        folder = _folder_for(relative_path)
        if folder:
            images_by_folder.setdefault(folder.casefold(), []).append((filename, content, _is_sku_path(relative_path)))

    product_codes = {_text(row["产品编号"]).casefold(): _text(row["产品编号"]) for row in product_rows}
    extra_folders = sorted(folder for folder in images_by_folder if folder not in product_codes)
    report_products = []
    imported_ids = []
    incoming_skus: dict[str, list[str]] = {}
    if not legacy:
        for row in product_rows:
            code = _text(row["产品编号"])
            entries = images_by_folder.get(code.casefold(), [])
            has_sku_folder = any(is_sku for _filename, _content, is_sku in entries)
            for filename, _content, is_sku in entries:
                if has_sku_folder and not is_sku:
                    continue
                stem = PurePosixPath(filename).stem.strip()
                if stem and not _pick_role(stem):
                    incoming_skus.setdefault(stem.casefold(), []).append(code)

    for row in product_rows:
        code = _text(row["产品编号"])
        files = images_by_folder.get(code.casefold(), [])
        has_sku_folder = any(is_sku for _filename, _content, is_sku in files)
        main_files = [(filename, content) for filename, content, is_sku in files if not is_sku]
        sku_files = [(filename, content) for filename, content, is_sku in files if is_sku]
        rows = [item for item in sku_rows if _text(item.get("产品编号")).casefold() == code.casefold()]
        product = db.scalar(select(ProductMaster).where(ProductMaster.spu_code == code))
        title = _text(row.get("商品标题")) or _text(row.get("商品名称")) or _text(row.get("英文名称")) or code
        if not product:
            product = ProductMaster(spu_code=code, title=title, category=CATEGORY)
            db.add(product)
            db.flush()
        product.title = title
        product.category = CATEGORY
        product.material = _text(row.get("主体材质 Main Material")) or _text(row.get("主体材质")) or None
        product.import_parameters_json = json.dumps(_normalized_product_parameters(row), ensure_ascii=False, default=str)
        product.image_rights = "AUTHORIZED"

        conflicts: list[str] = []
        if not legacy:
            rows = []
            sku_candidates = sku_files if has_sku_folder else [(filename, content) for filename, content, _is_sku in files if not _pick_role(PurePosixPath(filename).stem)]
            for filename, _content in sku_candidates:
                stem = PurePosixPath(filename).stem.strip()
                if not stem or _pick_role(stem):
                    continue
                existing_owner = db.scalar(select(SkuVariant.product_id).where(func.lower(SkuVariant.sku_code) == stem.casefold()).limit(1))
                if len(set(incoming_skus.get(stem.casefold(), []))) > 1 or (existing_owner is not None and existing_owner != product.id):
                    conflicts.append(stem)
                    continue
                rows.append({
                    "SKU编号": stem,
                    "颜色": stem,
                    "尺码/规格": row.get("尺码/规格"),
                    "SKU内单品件数": row.get("SKU内单品件数"),
                    "申报价(CNY)": row.get("供货价(CNY)") if row.get("供货价(CNY)") not in (None, "") else row.get("申报价(CNY)"),
                    "库存数量": row.get("库存数量"),
                })

        sku_map: dict[str, SkuVariant] = {}
        for index, sku_row in enumerate(rows, 1):
            sku_code = _text(sku_row.get("SKU编号")) or _text(sku_row.get("SKU货号")) or f"{code}-SKU-{index:02d}"
            sku = db.scalar(select(SkuVariant).where(SkuVariant.sku_code == sku_code))
            if not sku:
                sku = SkuVariant(product_id=product.id, sku_code=sku_code)
                db.add(sku)
            sku.color = _text(sku_row.get("颜色")) or "默认"
            sku.size = _text(sku_row.get("尺码/规格")) or None
            sku.quantity = max(1, int(_number(sku_row.get("SKU内单品件数"), 1)))
            sku.price = _number(sku_row.get("申报价(CNY)"), 0)
            sku.stock = max(0, int(_number(sku_row.get("库存数量"), 0)))
            sku.name = " / ".join(filter(None, [sku.color, sku.size]))
            sku.status = "READY"
            sku.is_sellable = True
            db.flush()
            sku_map[sku_code.casefold()] = sku

        # A fresh finished-image import defines the current sellable SKU set.
        # Retain historical rows for draft references, but exclude removed SKUs
        # from image-count checks and future marketplace payloads.
        current_sku_ids = {sku.id for sku in sku_map.values()}
        existing_skus = db.scalars(
            select(SkuVariant).where(SkuVariant.product_id == product.id)
        ).all()
        for existing_sku in existing_skus:
            existing_sku.is_sellable = existing_sku.id in current_sku_ids

        if rows:
            product.price = next((sku.price for sku in sku_map.values() if sku.price), 0)
            product.stock = sum(sku.stock for sku in sku_map.values())
            dims = [_text(_first_value(row, new_name, old_name, old_name.replace("(cm)", "（cm）"))) for new_name, old_name in (
                ("包装最长边 CM(批量)", "包装最长边(cm)"),
                ("包装次长边 CM(批量)", "包装次长边(cm)"),
                ("包装最短边 CM(批量)", "包装最短边(cm)"),
            )]
            product.dimensions = " x ".join(item for item in dims if item) or "见SKU规格"
            product.weight_g = _weight_in_grams(row, ("重量 KG(批量)", "商品净重 KG(批量)"), ("商品净重(g)", "商品净重（g）")) or None

        fingerprint = hashlib.sha256((code + "|" + "|".join(f"{name}:{hashlib.sha256(content).hexdigest()}" for name, content, _is_sku in files)).encode()).hexdigest()
        batch = ImageBatch(
            product_id=product.id,
            input_fingerprint=fingerprint,
            pipeline_version="finished-upload-v1",
            pipeline_kind="FINISHED_UPLOAD",
            provider="MANUAL_UPLOAD",
            version=(db.scalar(select(func.max(ImageBatch.version)).where(ImageBatch.product_id == product.id, ImageBatch.pipeline_kind == "FINISHED_UPLOAD")) or 0) + 1,
            status="NEEDS_ATTENTION",
            total_jobs=6 + len(sku_map),
        )
        db.add(batch)
        db.flush()

        assigned: dict[str, str] = {}
        remaining = []
        if has_sku_folder:
            for filename, content in sku_files:
                stem = PurePosixPath(filename).stem
                sku = sku_map.get(stem.casefold()) if not legacy else next((item for sku_code, item in sku_map.items() if _norm(sku_code) in _norm(stem)), None)
                if sku:
                    _save_finished_asset(db, product, batch, "SKU_WHITE", content, filename, sku)
                    assigned[f"SKU:{sku.sku_code}"] = filename
                else:
                    remaining.append((filename, content))
            for role, (filename, content) in zip(MAIN_ROLES, main_files):
                _save_finished_asset(db, product, batch, role, content, filename)
                assigned[role] = filename
            remaining.extend(main_files[len(MAIN_ROLES):])
        else:
            for filename, content, _is_sku in files:
                stem = PurePosixPath(filename).stem
                sku = sku_map.get(stem.casefold()) if not legacy else next((item for sku_code, item in sku_map.items() if _norm(sku_code) in _norm(stem)), None)
                role = "SKU_WHITE" if sku else _pick_role(stem)
                if role and (role == "SKU_WHITE" or role not in assigned):
                    _save_finished_asset(db, product, batch, role, content, filename, sku)
                    assigned[f"SKU:{sku.sku_code}" if sku else role] = filename
                else:
                    remaining.append((filename, content))

        if legacy:
            for role in MAIN_ROLES:
                if role not in assigned and remaining:
                    filename, content = remaining.pop(0)
                    _save_finished_asset(db, product, batch, role, content, filename)
                    assigned[role] = filename
            for sku in sku_map.values():
                key = f"SKU:{sku.sku_code}"
                if key not in assigned and remaining:
                    filename, content = remaining.pop(0)
                    _save_finished_asset(db, product, batch, "SKU_WHITE", content, filename, sku)
                    assigned[key] = filename

        missing = [role for role in MAIN_ROLES if role not in assigned]
        missing += [f"SKU:{sku.sku_code}" for sku in sku_map.values() if f"SKU:{sku.sku_code}" not in assigned]
        if not sku_map:
            missing.append("NO_SKU_IMAGES")
        batch.completed_jobs = len(assigned)
        batch.failed_jobs = 0
        batch.waiting_jobs = len(missing)
        batch.status = "COMPLETED" if not missing else "NEEDS_ATTENTION"
        batch.completed_at = datetime.now(timezone.utc)
        product.image_readiness = "GENERATED" if not missing else "NEEDS_ATTENTION"
        product.scene_readiness = "READY_FOR_REVIEW" if all(role in assigned for role in ("SCENE_MODEL_WEAR", "SCENE_LIFESTYLE")) else "NEEDS_ATTENTION"
        product.status = "READY" if not missing and not conflicts else "WAITING_GENERATION"
        imported_ids.append(product.id)
        report_products.append({"product_id": product.id, "product_code": code, "folder_found": bool(files), "image_count": len(files), "assigned": assigned, "missing": missing, "conflicts": sorted(set(conflicts)), "unused_files": [name for name, _ in remaining], "detected_skus": [{"sku_code": sku.sku_code, "name": sku.name, "image": assigned.get(f"SKU:{sku.sku_code}")} for sku in sku_map.values()]})

    db.commit()
    complete = lambda item: not item["missing"] and not item["conflicts"]
    return {"status": "READY_FOR_REVIEW" if all(complete(item) for item in report_products) else "NEEDS_ATTENTION", "product_count": len(report_products), "matched_count": sum(complete(item) for item in report_products), "extra_folders": extra_folders, "products": report_products, "product_ids": imported_ids}
