from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import (
    Asset,
    ImportBatch,
    ImportIssue,
    ProductMaster,
    SkuComponent,
    SkuVariant,
    Store,
    StoreListing,
)


def norm(raw: object) -> str:
    text = "" if raw is None else str(raw).strip()
    return re.sub(r"[\s\-_/（）()]+", "", text).lower()


ALIASES: Dict[str, Tuple[str, ...]] = {
    "spu_code": ("spu", "spucode", "spu编码", "商品编码", "产品编码", "货号"),
    "sku_code": ("sku", "skucode", "sku编码", "变体编码", "规格编码"),
    "title": ("title", "name", "商品标题", "商品名称", "产品名称", "标题", "sku名称"),
    "category": ("category", "类目", "商品类目", "分类"),
    "material": ("material", "材质", "材料"),
    "color": ("color", "颜色", "色"),
    "size": ("size", "尺寸", "规格", "长度"),
    "dimensions": ("dimensions", "商品尺寸", "包装尺寸"),
    "weight_g": ("weight", "weightg", "重量", "重量g", "克重"),
    "cost": ("cost", "成本", "采购价"),
    "price": ("price", "供货价", "售价", "价格", "零售价"),
    "currency": ("currency", "币种", "货币"),
    "stock": ("stock", "库存", "可售库存"),
    "quantity": ("quantity", "qty", "数量", "件数", "套装数量"),
    "source_image_url": ("image", "imageurl", "图片", "图片链接", "主图", "原图", "图片地址"),
    "image_rights": ("imagerights", "图片授权", "授权状态", "图片版权"),
    "is_sellable": ("issellable", "可售", "是否可售", "启用"),
    "store_name": ("store", "storename", "店铺", "店铺名称"),
    "platform": ("platform", "平台"),
    "mode": ("mode", "经营模式", "托管模式"),
    "component_code": ("componentcode", "组件编码", "子件编码"),
    "component_name": ("componentname", "组件名称", "组合商品", "子件名称"),
    "component_color": ("componentcolor", "组件颜色", "子件颜色"),
    "component_quantity": ("componentquantity", "组件数量", "子件数量", "组合数量"),
}


@dataclass
class SheetRows:
    name: str
    headers: Dict[str, str]
    rows: List[Tuple[int, Dict[str, object]]]


def header_map(values: Iterable[object]) -> Dict[str, str]:
    normalized = {norm(item): str(item).strip() for item in values if item is not None}
    result: Dict[str, str] = {}
    for canonical, aliases in ALIASES.items():
        for alias in (canonical, *aliases):
            if norm(alias) in normalized:
                result[canonical] = normalized[norm(alias)]
                break
    return result


def read_sheet(ws) -> SheetRows:
    iterator = ws.iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration:
        return SheetRows(ws.title, {}, [])
    mapping = header_map(raw_headers)
    rows: List[Tuple[int, Dict[str, object]]] = []
    for row_number, values in enumerate(iterator, start=2):
        raw = {
            str(raw_headers[index]).strip(): item
            for index, item in enumerate(values)
            if index < len(raw_headers) and raw_headers[index] is not None
        }
        if any(item not in (None, "") for item in raw.values()):
            rows.append((row_number, raw))
    return SheetRows(ws.title, mapping, rows)


def get_value(sheet: SheetRows, row: Dict[str, object], field: str, default: object = None) -> object:
    column = sheet.headers.get(field)
    return row.get(column, default) if column else default


def as_text(raw: object) -> Optional[str]:
    if raw is None:
        return None
    text = str(raw).strip()
    return text or None


def as_float(raw: object) -> Optional[float]:
    if raw in (None, ""):
        return None
    try:
        return float(str(raw).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def as_int(raw: object, default: int = 0) -> int:
    value = as_float(raw)
    return int(value) if value is not None else default


def as_bool(raw: object, default: bool = True) -> bool:
    if raw in (None, ""):
        return default
    return norm(raw) not in {"0", "false", "否", "no", "n", "停用"}


def choose_sheet(
    sheets: Dict[str, SheetRows], candidates: Tuple[str, ...], fallback: Optional[SheetRows] = None
) -> Optional[SheetRows]:
    for key, sheet in sheets.items():
        if any(norm(candidate) in key for candidate in candidates):
            return sheet
    return fallback


def add_issue(
    db: Session,
    batch: ImportBatch,
    sheet: str,
    row_number: int,
    identifier: Optional[str],
    field: Optional[str],
    code: str,
    message: str,
    severity: str = "ERROR",
) -> None:
    db.add(
        ImportIssue(
            batch_id=batch.id,
            sheet=sheet,
            row_number=row_number,
            identifier=identifier,
            field=field,
            code=code,
            message=message,
            severity=severity,
        )
    )


def sync_asset(
    db: Session,
    product: ProductMaster,
    url: Optional[str],
    rights: str,
    asset_type: str,
    sku: Optional[SkuVariant] = None,
) -> None:
    if not url:
        return
    stmt = select(Asset).where(
        Asset.product_id == product.id,
        Asset.sku_id == (sku.id if sku else None),
        Asset.asset_type == asset_type,
    )
    asset = db.scalar(stmt)
    if not asset:
        asset = Asset(product_id=product.id, sku_id=sku.id if sku else None, asset_type=asset_type, url=url)
        db.add(asset)
    asset.url = url
    asset.source_url = url
    asset.role = "SPU_MAIN_SOURCE" if asset_type == "SPU_SOURCE" else asset_type
    asset.mirror_status = "PENDING"
    asset.rights_status = rights


def import_products(db: Session, batch: ImportBatch, sheet: Optional[SheetRows]) -> None:
    if not sheet or not sheet.rows:
        add_issue(db, batch, "Products", 1, None, None, "EMPTY_PRODUCTS", "没有找到商品数据行")
        return
    for row_number, row in sheet.rows:
        spu_code = as_text(get_value(sheet, row, "spu_code"))
        title = as_text(get_value(sheet, row, "title"))
        identifier = spu_code or f"第{row_number}行"
        missing = []
        for field, label in (
            ("spu_code", "SPU编码"),
            ("title", "商品标题"),
            ("category", "类目"),
            ("material", "材质"),
        ):
            if not as_text(get_value(sheet, row, field)):
                missing.append(label)
        if missing:
            add_issue(
                db,
                batch,
                sheet.name,
                row_number,
                identifier,
                None,
                "MISSING_REQUIRED",
                f"缺少必填字段：{'、'.join(missing)}",
            )
        source_image = as_text(get_value(sheet, row, "source_image_url"))
        rights = as_text(get_value(sheet, row, "image_rights")) or "UNKNOWN"
        if not source_image:
            add_issue(
                db,
                batch,
                sheet.name,
                row_number,
                identifier,
                "source_image_url",
                "MISSING_SOURCE_IMAGE",
                "缺少商品原图，需补充素材",
                "WARNING",
            )
        if rights.upper() in {"UNKNOWN", "未确认", "不清楚"}:
            add_issue(
                db,
                batch,
                sheet.name,
                row_number,
                identifier,
                "image_rights",
                "UNKNOWN_IMAGE_RIGHTS",
                "图片授权状态未确认",
                "WARNING",
            )
        if not spu_code:
            continue
        product = db.scalar(select(ProductMaster).where(ProductMaster.spu_code == spu_code))
        if not product:
            product = ProductMaster(spu_code=spu_code, title=title or "待补商品标题")
            db.add(product)
        product.title = title or product.title
        product.category = as_text(get_value(sheet, row, "category")) or product.category
        product.material = as_text(get_value(sheet, row, "material"))
        product.color = as_text(get_value(sheet, row, "color"))
        product.dimensions = as_text(get_value(sheet, row, "dimensions"))
        product.weight_g = as_float(get_value(sheet, row, "weight_g"))
        product.cost = as_float(get_value(sheet, row, "cost"))
        product.price = as_float(get_value(sheet, row, "price"))
        product.currency = as_text(get_value(sheet, row, "currency")) or "USD"
        product.stock = as_int(get_value(sheet, row, "stock"))
        product.source_image_url = source_image
        product.image_rights = rights
        product.imported_batch_id = batch.id
        product.status = (
            "WAITING_DATA"
            if missing or not source_image or rights.upper() in {"UNKNOWN", "未确认", "不清楚"}
            else "WAITING_GENERATION"
        )
        db.flush()
        sync_asset(db, product, source_image, rights, "SPU_SOURCE")


def import_skus(db: Session, batch: ImportBatch, sheet: Optional[SheetRows]) -> None:
    if not sheet:
        add_issue(db, batch, "SKUs", 1, None, None, "MISSING_SKU_SHEET", "没有找到 SKU 工作表")
        return
    for row_number, row in sheet.rows:
        spu_code = as_text(get_value(sheet, row, "spu_code"))
        sku_code = as_text(get_value(sheet, row, "sku_code"))
        identifier = sku_code or f"第{row_number}行"
        if not spu_code or not sku_code:
            add_issue(db, batch, sheet.name, row_number, identifier, None, "MISSING_SKU_ID", "SKU 行必须包含 SPU 编码和 SKU 编码")
            continue
        product = db.scalar(select(ProductMaster).where(ProductMaster.spu_code == spu_code))
        if not product:
            add_issue(db, batch, sheet.name, row_number, sku_code, "spu_code", "UNKNOWN_SPU", f"SKU 对应的 SPU 不存在：{spu_code}")
            continue
        sku = db.scalar(select(SkuVariant).where(SkuVariant.sku_code == sku_code))
        if not sku:
            sku = SkuVariant(sku_code=sku_code, product_id=product.id)
            db.add(sku)
        sku.product_id = product.id
        sku.name = as_text(get_value(sheet, row, "title")) or sku.name
        sku.color = as_text(get_value(sheet, row, "color"))
        sku.size = as_text(get_value(sheet, row, "size"))
        sku.material = as_text(get_value(sheet, row, "material")) or product.material
        sku.quantity = as_int(get_value(sheet, row, "quantity"), 1)
        sku.price = as_float(get_value(sheet, row, "price")) or product.price
        sku.stock = as_int(get_value(sheet, row, "stock"))
        sku.source_image_url = as_text(get_value(sheet, row, "source_image_url"))
        sku.is_sellable = as_bool(get_value(sheet, row, "is_sellable"), True)
        sku.status = "WAITING_DATA" if not sku.source_image_url or not sku.color else "WAITING_GENERATION"
        db.flush()
        sync_asset(db, product, sku.source_image_url, product.image_rights, "SKU_SOURCE", sku)


def import_components(db: Session, batch: ImportBatch, sheet: Optional[SheetRows]) -> None:
    if not sheet:
        return
    cleared: set[int] = set()
    for row_number, row in sheet.rows:
        parent_code = as_text(get_value(sheet, row, "sku_code"))
        component_name = as_text(get_value(sheet, row, "component_name"))
        if not parent_code or not component_name:
            add_issue(db, batch, sheet.name, row_number, parent_code, None, "MISSING_COMPONENT", "组合行必须包含父 SKU 和组件名称")
            continue
        sku = db.scalar(select(SkuVariant).where(SkuVariant.sku_code == parent_code))
        if not sku:
            add_issue(db, batch, sheet.name, row_number, parent_code, "sku_code", "UNKNOWN_PARENT_SKU", f"组合父 SKU 不存在：{parent_code}")
            continue
        if sku.id not in cleared:
            for component in list(sku.components):
                db.delete(component)
            cleared.add(sku.id)
        component_image = as_text(get_value(sheet, row, "source_image_url"))
        db.add(
            SkuComponent(
                sku_id=sku.id,
                component_code=as_text(get_value(sheet, row, "component_code")),
                component_name=component_name,
                color=as_text(get_value(sheet, row, "component_color")) or as_text(get_value(sheet, row, "color")),
                quantity=as_int(get_value(sheet, row, "component_quantity"), 1),
                source_image_url=component_image,
            )
        )


def import_listings(db: Session, batch: ImportBatch, sheet: Optional[SheetRows]) -> None:
    if not sheet:
        return
    for row_number, row in sheet.rows:
        spu_code = as_text(get_value(sheet, row, "spu_code"))
        store_name = as_text(get_value(sheet, row, "store_name"))
        platform = as_text(get_value(sheet, row, "platform"))
        if not spu_code or not store_name or not platform:
            add_issue(db, batch, sheet.name, row_number, spu_code, None, "MISSING_STORE_TARGET", "店铺版本必须包含 SPU、店铺和平台")
            continue
        product = db.scalar(select(ProductMaster).where(ProductMaster.spu_code == spu_code))
        if not product:
            add_issue(db, batch, sheet.name, row_number, spu_code, "spu_code", "UNKNOWN_LISTING_SPU", f"店铺版本的 SPU 不存在：{spu_code}")
            continue
        store = db.scalar(select(Store).where(Store.name == store_name))
        if not store:
            store = Store(
                name=store_name,
                platform=platform,
                mode=as_text(get_value(sheet, row, "mode")) or "POP",
                currency=as_text(get_value(sheet, row, "currency")) or product.currency,
            )
            db.add(store)
            db.flush()
        listing = db.scalar(
            select(StoreListing).where(
                StoreListing.product_id == product.id, StoreListing.store_id == store.id
            )
        )
        if not listing:
            listing = StoreListing(product_id=product.id, store_id=store.id)
            db.add(listing)
        listing.listing_title = as_text(get_value(sheet, row, "title")) or product.title
        listing.price = as_float(get_value(sheet, row, "price")) or product.price
        listing.status = "NOT_READY"


def import_workbook(db: Session, content: bytes, filename: str) -> ImportBatch:
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheets = {norm(ws.title): read_sheet(ws) for ws in workbook.worksheets}
    fallback = next(iter(sheets.values()), None)
    product_sheet = choose_sheet(sheets, ("product", "products", "商品", "spu"), fallback)
    sku_sheet = choose_sheet(sheets, ("sku", "skus", "变体", "规格"))
    component_sheet = choose_sheet(sheets, ("component", "组合", "组件", "套装"))
    listing_sheet = choose_sheet(sheets, ("store", "listing", "店铺", "渠道"))

    batch = ImportBatch(filename=filename, status="PROCESSING")
    db.add(batch)
    db.flush()
    batch.total_rows = sum(len(sheet.rows) for sheet in sheets.values())

    import_products(db, batch, product_sheet)
    db.flush()
    import_skus(db, batch, sku_sheet)
    db.flush()
    import_components(db, batch, component_sheet)
    db.flush()
    for product in db.scalars(
        select(ProductMaster).where(ProductMaster.imported_batch_id == batch.id)
    ).all():
        for sku in product.skus:
            if not sku.is_sellable:
                continue
            component_ready = bool(sku.components) and all(
                component.source_image_url and component.quantity > 0 for component in sku.components
            )
            source_ready = bool(sku.source_image_url) or component_ready
            sku.status = "WAITING_GENERATION" if source_ready and sku.color else "WAITING_DATA"
            if not source_ready:
                add_issue(
                    db,
                    batch,
                    "SKUs",
                    0,
                    sku.sku_code,
                    "source_image_url",
                    "MISSING_SKU_IMAGE",
                    "缺少 SKU 原图或完整套装组件原图，禁止 AI 猜测生成",
                )
    import_listings(db, batch, listing_sheet)
    db.flush()

    for product in db.scalars(
        select(ProductMaster).where(ProductMaster.imported_batch_id == batch.id)
    ).all():
        if product.status == "WAITING_GENERATION" and (
            not product.skus
            or any(sku.status == "WAITING_DATA" for sku in product.skus if sku.is_sellable)
        ):
            product.status = "WAITING_DATA"

    db.flush()
    # SQLite and PostgreSQL both expose the relationship after flush; using len keeps this adapter portable.
    issue_count = len(batch.issues)
    batch.issue_rows = issue_count
    batch.valid_rows = max(batch.total_rows - issue_count, 0)
    batch.status = "COMPLETED_WITH_ISSUES" if issue_count else "COMPLETED"
    db.commit()
    db.refresh(batch)
    return batch


def template_workbook() -> Workbook:
    wb = Workbook()
    products = wb.active
    products.title = "Products"
    products.append(["spu_code", "title", "category", "material", "color", "dimensions", "weight_g", "cost", "price", "currency", "stock", "source_image_url", "image_rights"])
    products.append(["SPU-DEMO-001", "Sample silver pendant set", "女士饰品/项链", "925银", "银色", "45cm", 12.6, 3.8, 16.8, "USD", 100, "https://example.com/product-source.png", "AUTHORIZED"])

    skus = wb.create_sheet("SKUs")
    skus.append(["spu_code", "sku_code", "title", "color", "size", "material", "quantity", "price", "stock", "source_image_url", "is_sellable"])
    skus.append(["SPU-DEMO-001", "SKU-DEMO-001-SILVER", "银色项链 / 45cm", "银色", "45cm", "925银", 1, 9.9, 100, "https://example.com/sku-necklace.png", True])
    skus.append(["SPU-DEMO-001", "SKU-DEMO-SET", "银色项链耳钉两件套", "银色", "45cm", "925银", 2, 16.8, 60, "https://example.com/sku-set.png", True])

    components = wb.create_sheet("Components")
    components.append(["sku_code", "component_code", "component_name", "component_color", "component_quantity", "source_image_url"])
    components.append(["SKU-DEMO-SET", "SKU-DEMO-001-SILVER", "银色项链", "银色", 1, "https://example.com/sku-necklace.png"])
    components.append(["SKU-DEMO-SET", "SKU-DEMO-EARRING", "银色耳钉", "银色", 1, "https://example.com/component-earring.png"])

    listings = wb.create_sheet("StoreListings")
    listings.append(["spu_code", "store_name", "platform", "mode", "title", "price", "currency"])
    listings.append(["SPU-DEMO-001", "环球饰品（深圳）", "TEMU", "全托管", "925 Silver Pendant Necklace Set", 16.8, "USD"])
    return wb
