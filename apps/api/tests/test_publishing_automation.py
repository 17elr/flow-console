from __future__ import annotations

import io
import json
import zipfile

from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.automation import recover_interrupted_runs, run_automation
from app.db import Base
from app.models import AutomationRun, ListingCopy, MiaoshouDraft, PackagingPreset, ProductMaster, ProductPackagingSelection, SkuVariant, Store, StoreAutomationConfig, StoreListing
from app.publishing import MIAOSHOU_FORMAT2_HEADERS, build_listing_workbook, create_aliexpress_import_packages, create_store_drafts, listing_package


def setup_product(session: Session) -> tuple[ProductMaster, Store]:
    store = Store(name="Ali Test", platform="AliExpress", mode="POP", currency="USD", active=True)
    product = ProductMaster(spu_code="PUB-001", title="Test necklace", category="Necklace", dimensions="45cm", price=10.5, stock=12, currency="USD", import_parameters_json=json.dumps({"产地（国家或地区） Origin": "中国大陆(Origin)", "高关注化学品 High-concerned chemical": "无"}, ensure_ascii=False))
    session.add_all([store, product]); session.flush()
    sku = SkuVariant(product_id=product.id, sku_code="PUB-001-S", name="Silver", color="Silver", size="45cm", quantity=1, price=11.0, stock=12, is_sellable=True)
    session.add(sku); session.flush()
    session.add(ListingCopy(product_id=product.id, platform="ALIEXPRESS", title="Silver Necklace 45cm", bullet_points_json='["Color: Silver."]', description="A necklace offered in the listed SKU.", sku_names_json=f'{{"{sku.id}": "Silver / 45cm"}}', input_fingerprint="c" * 64, status="REVIEWED", issues_json="[]"))
    session.add(StoreAutomationConfig(store_id=store.id, price_multiplier=1.2, visual_profile="neutral-commerce", max_daily=20))
    session.commit(); session.refresh(product); session.refresh(store)
    return product, store


def test_excel_import_package_has_typed_product_and_sku_mapping() -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        product, store = setup_product(session)
        content = build_listing_workbook(session, product, store)
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        assert workbook.sheetnames == ["Product", "SKUs", "Image Mapping", "Category Parameters", "Instructions"]
        assert workbook["Product"]["A2"].value == "PUB-001"
        assert workbook["Product"]["F2"].value == 12.6
        assert isinstance(workbook["Product"]["F2"].value, float)
        assert workbook["SKUs"]["A2"].value == "PUB-001-S"
        assert workbook["SKUs"]["F2"].value == 13.2
        assert workbook["SKUs"]["H2"].value is None
        assert workbook["Product"].freeze_panes == "A2"


def test_automation_skips_unapproved_products_without_failing_run() -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        _, store = setup_product(session)
        run = run_automation(session, [store.id], 20)
        assert run.status == "COMPLETED"
        assert run.processed == 0
        assert run.skipped == 1
        assert run.failed == 0
        assert "等待人工审核" in run.details_json


def test_recover_interrupted_runs_closes_stale_running_records() -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        run = AutomationRun(trigger="MANUAL", status="RUNNING", details_json="[]")
        session.add(run); session.commit()
        assert recover_interrupted_runs(session) == 1
        session.refresh(run)
        assert run.status == "INTERRUPTED"
        assert run.failed == 1
        assert "API 进程重启" in run.details_json
        assert run.completed_at is not None


def test_automation_does_not_count_publishing_as_failure(monkeypatch) -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    monkeypatch.setattr("app.automation.workflow_payload", lambda *_args: {"review": {"approved": True}})
    monkeypatch.setattr("app.automation.create_store_drafts", lambda *_args, **_kwargs: [{"store_id": 1, "status": "PUBLISHING"}])
    with Session(engine) as session:
        product, store = setup_product(session)
        run = run_automation(session, [store.id], 20)
        assert run.status == "WAITING_EXTERNAL"
        assert run.processed == 1
        assert run.succeeded == 0
        assert run.failed == 0


def test_temu_create_uses_excel_product_title(monkeypatch) -> None:
    # The workbook's 商品名称 is authoritative for Miaoshou. ListingCopy.title
    # is review copy and must not replace the imported product title.
    source = __import__("inspect").getsource(create_store_drafts)
    assert '"title": excel_title' in source
    assert 'approved_title' not in source


def test_listing_package_contains_selected_packaging_image(monkeypatch) -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)

    image_stream = io.BytesIO()
    Image.new("RGB", (20, 20), "white").save(image_stream, "PNG")

    class Storage:
        def get(self, key: str) -> bytes:
            assert key == "packaging/preset-1.png"
            return image_stream.getvalue()

    monkeypatch.setattr("app.publishing.configured_storage", lambda: Storage())
    with Session(engine) as session:
        product, store = setup_product(session)
        preset = PackagingPreset(slot=1, storage_key="packaging/preset-1.png", sha256="a" * 64, mime_type="image/png", width=900, height=900, byte_size=13, original_filename="package.png")
        session.add(preset); session.flush()
        session.add(ProductPackagingSelection(product_id=product.id, preset_id=preset.id)); session.commit()
        package, _ = listing_package(session, product, store)
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            root = "产品素材包模版2/本地导入素材/"
            workbook_path = f"{root}产品导入表格.xlsx"
            packaging_path = f"{root}产品图片/PUB-001/详情图/详情图_7.jpg"
            assert archive.testzip() is None
            assert workbook_path in archive.namelist()
            assert packaging_path in archive.namelist()
            with Image.open(io.BytesIO(archive.read(packaging_path))) as image:
                assert image.format == "JPEG"
            workbook = load_workbook(io.BytesIO(archive.read(workbook_path)))
            sheet = workbook.worksheets[0]
            expected_headers = list(MIAOSHOU_FORMAT2_HEADERS)
            expected_headers[7:9] = ["规格1（颜色）", "SKU规格2（尺寸）"]
            assert [cell.value for cell in sheet[2]] == expected_headers
            assert sheet.max_row == 3
            assert sheet["A3"].value == "PUB-001"
            assert sheet["D3"].value == "PUB-001"
            assert sheet["E3"].value == "手动创建"
            assert sheet["H3"].value == "Silver"
            assert sheet["J3"].value == "PUB-001-S"
            assert sheet["C3"].value == "CNY"
            assert sheet["K3"].value == 11.0
            assert sheet["F3"].value == "A necklace offered in the listed SKU."
            assert sheet["G3"].value is None


def test_aliexpress_import_package_is_local_and_idempotent(monkeypatch) -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)

    class Storage:
        def __init__(self):
            self.items = {}

        def get(self, key: str) -> bytes:
            return self.items[key]

        def put(self, key: str, content: bytes, _content_type: str) -> None:
            self.items[key] = content

    storage = Storage()
    monkeypatch.setattr("app.publishing.configured_storage", lambda: storage)
    with Session(engine) as session:
        product, store = setup_product(session)
        state = {"review": {"fingerprint": "approved"}}
        result = create_aliexpress_import_packages(session, product, state, [store.id])
        assert result[0]["status"] == "PACKAGE_READY"
        assert result[0]["package_available"] is True
        draft = session.scalar(select(MiaoshouDraft).where(MiaoshouDraft.store_id == store.id))
        assert draft is not None and draft.package_key
        with zipfile.ZipFile(io.BytesIO(storage.get(draft.package_key))) as archive:
            workbook_path = "产品素材包模版2/本地导入素材/产品导入表格.xlsx"
            assert archive.testzip() is None
            assert workbook_path in archive.namelist()
            assert load_workbook(io.BytesIO(archive.read(workbook_path))).sheetnames == ["Sheet1", "Sheet2", "Sheet3"]
        again = create_aliexpress_import_packages(session, product, state, [store.id])
        assert again[0]["idempotent"] is True


def test_historical_published_record_is_downgraded_until_miaoshou_verifies(monkeypatch) -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)

    class Provider:
        calls = 0

        def wait_for_publish(self, detail_id, mode, **_kwargs):
            assert detail_id == "456"
            self.calls += 1
            return {"state": "notPublished", "verified": False, "checked_at": "2026-08-23T00:00:00+00:00"}

        def ensure_shop_item_number(self, shop_id, detail_id, cid, item_number, mode, *_args):
            assert (shop_id, detail_id, item_number) == ("123", "456", "PUB-001")
            return {"updated": True, "response": {"result": "success"}}

        def publish_product(self, shop_id, detail_id, mode):
            assert (shop_id, detail_id) == ("123", "456")
            return {"response": {"result": "success"}}

    monkeypatch.setattr("app.publishing.configured_miaoshou", lambda: Provider())
    monkeypatch.setattr("app.publishing.publishing_key", lambda *_args, **_kwargs: "existing-key")
    with Session(engine) as session:
        product, old_store = setup_product(session)
        session.delete(old_store)
        store = Store(name="TEMU Test", platform="TEMU", mode="FULL", currency="USD", active=True, external_shop_id="123")
        session.add(store); session.flush()
        session.add(ListingCopy(product_id=product.id, platform="TEMU", title="Pendant", bullet_points_json="[]", description="Pendant", sku_names_json="{}", input_fingerprint="d" * 64, status="REVIEWED", issues_json="[]"))
        listing = StoreListing(product_id=product.id, store_id=store.id, status="PUBLISHED", external_product_id="456")
        draft = MiaoshouDraft(product_id=product.id, store_id=store.id, idempotency_key="existing-key", channel="TEMU", status="PUBLISHED", external_id="456", response_json='{"publish":{"result":"success"}}')
        session.add_all([listing, draft]); session.commit(); session.refresh(product)

        result = create_store_drafts(session, product, {"review": {"fingerprint": "approved"}}, [store.id], auto_publish=True)
        session.refresh(draft); session.refresh(listing)
        assert result[0]["status"] == "PUBLISHING"
        assert result[0]["verification_required"] is True
        assert result[0]["external_id"] == "456"
        assert draft.status == "PUBLISHING"
        assert listing.status == "PUBLISHING"
        assert '"state": "notPublished"' in draft.response_json
        assert '"item_number_repair"' in draft.response_json
        assert '"publish_retry"' in draft.response_json


def test_historical_published_record_stays_published_only_after_verification(monkeypatch) -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)

    class Provider:
        def wait_for_publish(self, _detail_id, _mode, **_kwargs):
            return {"state": "published", "verified": True, "checked_at": "2026-08-23T00:00:00+00:00"}

    monkeypatch.setattr("app.publishing.configured_miaoshou", lambda: Provider())
    monkeypatch.setattr("app.publishing.publishing_key", lambda *_args, **_kwargs: "existing-key")
    with Session(engine) as session:
        product, old_store = setup_product(session)
        session.delete(old_store)
        store = Store(name="TEMU Test", platform="TEMU", mode="FULL", currency="USD", active=True, external_shop_id="123")
        session.add(store); session.flush()
        session.add(ListingCopy(product_id=product.id, platform="TEMU", title="Pendant", bullet_points_json="[]", description="Pendant", sku_names_json="{}", input_fingerprint="d" * 64, status="REVIEWED", issues_json="[]"))
        draft = MiaoshouDraft(product_id=product.id, store_id=store.id, idempotency_key="existing-key", channel="TEMU", status="PUBLISHED", external_id="456")
        session.add(draft); session.commit(); session.refresh(product)

        result = create_store_drafts(session, product, {"review": {"fingerprint": "approved"}}, [store.id])
        assert result[0]["status"] == "PUBLISHED"
        assert result[0]["verification_required"] is False
