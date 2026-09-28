from __future__ import annotations

import json
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db import Base, get_db
from app.main import app
from app.models import ProductMaster
from app.finished_import import ALIEXPRESS_CATEGORY, _weight_in_grams


def test_aliexpress_category_matches_source_workbook() -> None:
    assert ALIEXPRESS_CATEGORY == "珠宝饰品及配件 (Jewelry & Accessories)/流行饰品 (Fashion Jewelry)/项链 (Necklace)"


class MemoryStorage:
    def __init__(self) -> None:
        self.items: dict[str, bytes] = {}

    def put(self, key: str, content: bytes, _content_type: str) -> None:
        self.items[key] = content

    def get(self, key: str) -> bytes:
        return self.items[key]


def test_merged_weight_column_uses_kilograms() -> None:
    row = {"重量 KG(批量)": 0.2, "商品净重 KG(批量)": 0.1}
    assert _weight_in_grams(row, ("重量 KG(批量)", "商品净重 KG(批量)"), ("商品净重(g)",)) == 200


def workbook_bytes() -> bytes:
    workbook = Workbook()
    product = workbook.active
    product.title = "商品参数"
    product.append(["title"])
    product.append(["note"])
    product.append(["产品编号", "商品名称", "主体材质 Main Material"])
    product.append(["P-100", "Heart necklace", "不锈钢"])
    sku = workbook.create_sheet("SKU规格")
    sku.append(["title"])
    sku.append(["note"])
    sku.append(["产品编号", "SKU编号", "颜色", "SKU内单品件数", "申报价(CNY)", "库存数量"])
    sku.append(["P-100", "P-100-GOLD", "金色", 1, 29.9, 20])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def single_sheet_workbook_bytes(
    products: list[tuple[str, str]],
    category_id: str | None = None,
    individually_packed: str = "不是独立包装",
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品参数"
    sheet.append(["title"])
    sheet.append(["note"])
    sheet.append(["group"])
    sheet.append(["产品编号", "妙手类目ID", "商品名称", "主体材质 Main Material", "尺码/规格", "SKU内单品件数", "申报价(CNY)", "库存数量", "包装最长边(cm)", "包装次长边(cm)", "包装最短边(cm)", "包装重量(g)", "商品净重(g)", "SKU是否独立包装"])
    for code, title in products:
        sheet.append([code, category_id, title, "不锈钢", "均码", 1, 29.9, 20, 12, 8, 2, 35, 12, individually_packed])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_finished_workbook_accepts_official_product_title_header() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品参数"
    sheet.append(["title"])
    sheet.append(["note"])
    sheet.append(["产品编号", "商品标题"])
    sheet.append(["P-TITLE", "Official title"])
    stream = BytesIO()
    workbook.save(stream)
    from app.finished_import import parse_finished_workbook
    products, skus, legacy = parse_finished_workbook(stream.getvalue())
    assert products == [{"产品编号": "P-TITLE", "商品标题": "Official title"}]
    assert skus == [] and legacy is False


def image_bytes(color: str = "white") -> bytes:
    stream = BytesIO()
    Image.new("RGB", (900, 900), color).save(stream, "PNG")
    return stream.getvalue()


def test_finished_folder_import_matches_product_and_skips_generation(monkeypatch) -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    storage = MemoryStorage()
    monkeypatch.setattr("app.finished_import.configured_storage", lambda: storage)

    def override_db():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    client = TestClient(app)
    files = [("files", (f"{index:02d}.png", image_bytes(), "image/png")) for index in range(1, 7)]
    files.append(("files", ("P-100-GOLD.png", image_bytes("gray"), "image/png")))
    paths = [("relative_paths", (None, f"成品图/P-100/{index:02d}.png")) for index in range(1, 7)]
    paths.append(("relative_paths", (None, "成品图/P-100/P-100-GOLD.png")))
    try:
        response = client.post(
            "/api/imports/finished-images",
            files=[("workbook", ("parameters.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")), *files, *paths],
        )
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["status"] == "READY_FOR_REVIEW"
        assert result["matched_count"] == 1
        assert result["products"][0]["missing"] == []
        assert result["products"][0]["assigned"]["SKU:P-100-GOLD"] == "P-100-GOLD.png"
        workflow = client.get(f"/api/products/{result['product_ids'][0]}/workflow").json()
        assert workflow["status"] == "READY_TO_PUBLISH"
        assert workflow["output_count"] == 7
        assert all(item["qc"]["status"] == "MANUAL_FINISHED_UPLOAD" for item in workflow["outputs"])
    finally:
        app.dependency_overrides.clear()


def test_single_sheet_auto_creates_skus_and_requires_packaging(monkeypatch) -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    storage = MemoryStorage()
    monkeypatch.setattr("app.finished_import.configured_storage", lambda: storage)
    monkeypatch.setattr("app.main.configured_storage", lambda: storage)

    def override_db():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    client = TestClient(app)
    uploads = [("files", (f"{index:02d}.png", image_bytes(), "image/png")) for index in range(1, 7)]
    uploads += [("files", ("UNIQUE-GOLD.png", image_bytes("yellow"), "image/png")), ("files", ("UNIQUE-SILVER.png", image_bytes("gray"), "image/png"))]
    paths = [("relative_paths", (None, f"成品图/P-200/{index:02d}.png")) for index in range(1, 7)]
    paths += [("relative_paths", (None, "成品图/P-200/UNIQUE-GOLD.png")), ("relative_paths", (None, "成品图/P-200/UNIQUE-SILVER.png"))]
    try:
        response = client.post("/api/imports/finished-images", files=[("workbook", ("v3.xlsx", single_sheet_workbook_bytes([("P-200", "Pendant")], "29542"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")), *uploads, *paths])
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["matched_count"] == 1
        assert [item["sku_code"] for item in result["products"][0]["detected_skus"]] == ["UNIQUE-GOLD", "UNIQUE-SILVER"]
        product_id = result["product_ids"][0]
        with Session(engine) as session:
            product = session.get(ProductMaster, product_id)
            assert json.loads(product.import_parameters_json)["cid"] == "29542"
            assert product.title == "Pendant"
            assert product.dimensions == "12 x 8 x 2"
            assert product.weight_g == 12
            assert json.loads(product.import_parameters_json)["individuallyPacked"] == 0
        workflow = client.get(f"/api/products/{product_id}/workflow").json()
        assert workflow["publish_blockers"] == ["请选择外包装图片"]
        preset = client.put("/api/packaging-presets/1", files={"file": ("package.png", image_bytes("blue"), "image/png")})
        assert preset.status_code == 200
        selected = client.put(f"/api/products/{product_id}/packaging-selection", json={"preset_id": preset.json()["id"]})
        assert selected.status_code == 200
        workflow = client.get(f"/api/products/{product_id}/workflow").json()
        assert workflow["publish_blockers"] == []
        assert workflow["packaging_selection"]["slot"] == 1
    finally:
        app.dependency_overrides.clear()


def test_single_sheet_preserves_main_upload_order_and_reads_sku_subfolder(monkeypatch) -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    storage = MemoryStorage()
    monkeypatch.setattr("app.finished_import.configured_storage", lambda: storage)

    def override_db():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    client = TestClient(app)
    main_names = ["cover.png", "detail-a.png", "gift-box.png", "measurements.png", "model.png", "occasion.png"]
    sku_names = ["SKU-GOLD.png", "SKU-SILVER.png"]
    uploads = [("files", (name, image_bytes(), "image/png")) for name in main_names + sku_names]
    paths = [("relative_paths", (None, f"成品图/P-ORDER/{name}")) for name in main_names]
    paths += [("relative_paths", (None, f"成品图/P-ORDER/sku/{name}")) for name in sku_names]
    try:
        response = client.post(
            "/api/imports/finished-images",
            files=[
                ("workbook", ("v8.xlsx", single_sheet_workbook_bytes([("P-ORDER", "Ordered pendant")]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
                *uploads,
                *paths,
            ],
        )
        assert response.status_code == 200, response.text
        product = response.json()["products"][0]
        assert [product["assigned"][role] for role in (
            "SPU_WHITE_MAIN", "SPU_DETAIL_1", "SPU_DETAIL_2", "SPU_SIZE_INFO", "SCENE_MODEL_WEAR", "SCENE_LIFESTYLE",
        )] == main_names
        assert [item["sku_code"] for item in product["detected_skus"]] == ["SKU-GOLD", "SKU-SILVER"]
        assert product["assigned"]["SKU:SKU-GOLD"] == "SKU-GOLD.png"
        assert product["assigned"]["SKU:SKU-SILVER"] == "SKU-SILVER.png"
        assert product["missing"] == []
    finally:
        app.dependency_overrides.clear()


def test_single_sheet_blocks_duplicate_sku_image_names(monkeypatch) -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    storage = MemoryStorage()
    monkeypatch.setattr("app.finished_import.configured_storage", lambda: storage)

    def override_db():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    client = TestClient(app)
    uploads = []
    paths = []
    for product in ("P-301", "P-302"):
        for index in range(1, 7):
            uploads.append(("files", (f"{index:02d}.png", image_bytes(), "image/png")))
            paths.append(("relative_paths", (None, f"成品图/{product}/{index:02d}.png")))
        uploads.append(("files", ("DUPLICATE-SKU.png", image_bytes("gray"), "image/png")))
        paths.append(("relative_paths", (None, f"成品图/{product}/DUPLICATE-SKU.png")))
    try:
        response = client.post("/api/imports/finished-images", files=[("workbook", ("v3.xlsx", single_sheet_workbook_bytes([("P-301", "One"), ("P-302", "Two")]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")), *uploads, *paths])
        assert response.status_code == 200, response.text
        assert response.json()["matched_count"] == 0
        assert all(item["conflicts"] == ["DUPLICATE-SKU"] for item in response.json()["products"])
    finally:
        app.dependency_overrides.clear()
