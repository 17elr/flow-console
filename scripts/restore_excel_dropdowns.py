from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "outputs" / "temu-template-v6" / "TEMU-女士时尚吊坠项链-单页产品参数模板-v6.xlsx"

wb = load_workbook(BOOK)
product = wb["商品参数"]
options = wb["选项表"]

# The v3 source workbook contains inline validations tied to its old column
# order. Keeping them causes WPS to open the wrong list even when a newer rule
# targets the same cell. Rebuild validation from zero against the current v6
# headers so adjacent attributes can never share each other's options.
product.data_validations.dataValidation = []

headers = {str(cell.value or "").strip(): cell.column for cell in product[4]}
option_headers = {str(cell.value or "").strip(): cell.column for cell in options[4]}

def col_letter(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result

for header, product_col in headers.items():
    if not header or header not in option_headers:
        continue
    option_col = option_headers[header]
    values = [options.cell(row, option_col).value for row in range(5, options.max_row + 1)]
    values = [value for value in values if value not in (None, "")]
    if not values:
        continue
    last_row = 4 + len(values)
    formula = f"'选项表'!${col_letter(option_col)}$5:${col_letter(option_col)}${last_row}"
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.errorTitle = "请选择列表中的值"
    validation.error = "多选字段可在同一单元格使用中文顿号“、”填写多个选项；需要全部选项时填写“全选”。"
    validation.showErrorMessage = False
    product.add_data_validation(validation)
    validation.add(f"{col_letter(product_col)}5:{col_letter(product_col)}200")

# "无" is not a valid birthstone value. Old template rows used it as a generic
# placeholder, which would be mistaken for a real selection after import.
birthstone_column = next(
    (column for header, column in headers.items() if "Birthstone" in header or "诞生石" in header),
    None,
)
if birthstone_column:
    valid_birthstones = {
        "全选", "一月生日石", "二月生日石", "三月生日石", "四月生日石",
        "五月生日石", "六月生日石", "七月生日石", "八月生日石",
        "九月生日石", "十月生日石", "十一月生日石", "十二月生日石",
    }
    for row in range(5, 201):
        cell = product.cell(row, birthstone_column)
        if cell.value and str(cell.value).strip() not in valid_birthstones:
            cell.value = None

# Fail template generation immediately if these adjacent fields are ever
# shifted again.
expected = {
    "Whether it contains metal components": {"是", "否"},
    "Birthstone": valid_birthstones,
    "Theme": {"全选", "动物", "体育", "假日", "食品和饮料", "航海主题", "天体符号", "脚本和数字符号", "鲜花", "户外标志", "心形", "宗教符号", "克拉达", "弓", "无穷大"},
}
for marker, allowed in expected.items():
    header = next((name for name in headers if marker in name), None)
    assert header, f"missing product column: {marker}"
    option_column = option_headers[header]
    actual = {
        str(options.cell(row, option_column).value).strip()
        for row in range(5, options.max_row + 1)
        if options.cell(row, option_column).value not in (None, "")
    }
    assert actual == allowed, f"dropdown mismatch for {marker}: {actual}"

wb.save(BOOK)
print(f"restored dropdowns in {BOOK}")
